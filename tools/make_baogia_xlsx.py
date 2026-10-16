"""Tạo submit/VBot_baogia.xlsx từ docs/baogia-A/B/C.tsv và evidence/INDEX.md.

Sheet BaoGia: 5 cột đúng đề, dòng cuối TOTAL (công thức SUM).
Sheet BangChung: evidence_id, mô tả, link, ngày, ảnh (ảnh chưa chụp: dán tay vào cột E).
Cột Bằng chứng được chuẩn hóa thành "E01; https://..." (mã ảnh và link, cách nhau dấu ;).

Chạy: python tools/make_baogia_xlsx.py submit/VBot_baogia.xlsx
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
out = Path(sys.argv[1] if len(sys.argv) > 1 else ROOT / "submit" / "VBot_baogia.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)

COMPONENTS = ("compute llm embedding gateway load_balancer authn authz database vector_db object_storage "
              "cache_queue sandbox connectors secrets_kms backup_dr monitoring governance network cicd "
              "ops_labor licenses contingency other").split()
NO_EVIDENCE = {"ops_labor", "contingency"}
URL_RE = re.compile(r"https?://[^\s;,)\]]+")
ID_RE = re.compile(r"\bE\d{2,3}\b")

# ---- sổ bằng chứng ----
index = []
for line in (ROOT / "evidence" / "INDEX.md").read_text(encoding="utf-8").splitlines():
    if not re.match(r"^\|\s*E\d{2,3}\s*\|", line):
        continue
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    code, desc = cells[0], cells[1]
    url = (URL_RE.search(cells[2]) or [None])[0] if len(cells) > 2 else None
    date = next((c for c in cells if re.fullmatch(r"\d{4}-\d{2}-\d{2}", c)), "")
    index.append((code, desc, url or "", date))
known = {c for c, *_ in index}


def norm_evidence(comp, text):
    if comp in NO_EVIDENCE:
        return "NA"
    text = text.replace(" ", " ")
    ids = []
    for m in ID_RE.findall(text):
        if m not in ids:
            ids.append(m)
    urls = []
    for m in URL_RE.findall(text):
        m = m.rstrip(".:")
        if m not in urls:
            urls.append(m)
    unknown = [i for i in ids if i not in known]
    if unknown:
        raise SystemExit(f"Mã bằng chứng không có trong evidence/INDEX.md: {unknown} (dòng {comp}: {text[:80]})")
    parts = ids + urls
    if not parts:
        raise SystemExit(f"Dòng {comp} không có mã ảnh hoặc link: {text[:80]}")
    return "; ".join(parts)


rows = []
for part in "ABC":
    with open(ROOT / "docs" / f"baogia-{part}.tsv", encoding="utf-8") as f:
        for r in list(csv.reader(f, delimiter="\t"))[1:]:
            comp, yn, gc, cost, ev = (c.strip() for c in r[:5])
            if comp not in COMPONENTS:
                raise SystemExit(f"Component lạ: {comp}")
            rows.append([comp, yn, gc, round(float(cost), 2), norm_evidence(comp, ev)])

# ---- workbook ----
wb = Workbook()
ws = wb.active
ws.title = "BaoGia"
header = ["Component", "Ý nghĩa", "Ghi chú", "Chi phí/tháng (USD)", "Bằng chứng"]
ws.append(header)
for r in rows:
    ws.append(r)
n = len(rows)
ws.append(["TOTAL", "", "", f"=SUM(D2:D{n + 1})", ""])

fill = PatternFill("solid", fgColor="D9EDF6")
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = fill
for row in ws.iter_rows(min_row=2, max_row=n + 2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    row[3].number_format = "#,##0.00"
for c in ws[n + 2]:
    c.font = Font(bold=True)
for col, w in zip("ABCDE", [16, 38, 70, 16, 52]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
dv = DataValidation(type="list", formula1='"' + ",".join(COMPONENTS) + '"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"A2:A{n + 1}")

eb = wb.create_sheet("BangChung")
eb.append(["evidence_id", "Mô tả", "Link", "Ngày chụp/tra", "Ảnh chụp màn hình (dán vào đây)"])
for c in eb[1]:
    c.font = Font(bold=True)
    c.fill = fill
for code, desc, url, date in index:
    eb.append([code, desc, url, date, "CHƯA CHỤP"])
for row in eb.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
for col, w in zip("ABCDE", [12, 80, 55, 14, 34]):
    eb.column_dimensions[col].width = w
eb.freeze_panes = "A2"

wb.save(out)
print(f"saved {out}: {n} dòng chi phí, {len(index)} mã bằng chứng")
